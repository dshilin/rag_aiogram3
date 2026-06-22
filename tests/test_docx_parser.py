from pathlib import Path
from docx import Document as DocxDocument

import pytest

from src.rag.docx_parser import DocxParser


def _make_docx(paragraphs: list[tuple], tmp_path: Path) -> Path:
    """Create a .docx for testing. Each tuple: (text, style) where style is 'h1'..'h4' or None for body."""
    doc = DocxDocument()
    for text, style in paragraphs:
        if style and style.startswith("h"):
            doc.add_heading(text, level=int(style[1]))
        else:
            doc.add_paragraph(text)
    path = tmp_path / "test.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def parser():
    return DocxParser()
